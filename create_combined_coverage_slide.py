from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


SOURCE = Path("/home/ubuntu/.cursor/projects/workspace/uploads/COverage_bfcb.pptx")
OUTPUT = Path("/workspace/COverage_combined.pptx")

NAVY = RGBColor(0x07, 0x18, 0x2D)
BLUE = RGBColor(0x0A, 0x60, 0xFF)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT_BLUE = RGBColor(0xE9, 0xF2, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def clone_shape(source_shape, destination_slide):
    """Copy a shape and recreate any image relationships on the destination slide."""
    element = deepcopy(source_shape.element)
    for node in element.iter():
        for attribute in (qn("r:embed"), qn("r:link")):
            old_rid = node.get(attribute)
            if not old_rid:
                continue
            related_part = source_shape.part.related_part(old_rid)
            new_rid = destination_slide.part.relate_to(related_part, RT.IMAGE)
            node.set(attribute, new_rid)
    destination_slide.shapes._spTree.insert_element_before(element, "p:extLst")
    return destination_slide.shapes[-1]


def add_text(slide, text, left, top, width, height, size, color, *, bold=False,
             align=PP_ALIGN.LEFT, font_name="CiscoSansTT", margin=0.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def direct_report_data(slide):
    """Match each report text box with the nearest circular headshot."""
    text_shapes = []
    headshots = []
    for shape in slide.shapes:
        if shape.name.startswith("Oval") and shape.fill.type is not None:
            # Exclude the much larger leader portrait.
            if shape.width < Inches(1.0):
                headshots.append(shape)
        if shape.has_text_frame:
            text = shape.text.strip().replace("\x0b", "\n")
            if "Account Executive" in text or "Solutions Engineer" in text:
                text_shapes.append(shape)

    pairs = []
    for text_shape in text_shapes:
        text_center = text_shape.left + text_shape.width / 2
        photo = min(
            headshots,
            key=lambda candidate: abs(
                candidate.left + candidate.width / 2 - text_center
            ),
        )
        lines = [line.strip() for line in text_shape.text.replace("\x0b", "\n").splitlines()]
        pairs.append((lines[0], lines[1], photo))
    return sorted(pairs, key=lambda item: item[2].left)


def add_report_row(slide, report, x, y, column_width, row_height):
    name, role, source_photo = report
    photo_size = Inches(0.43)
    photo = clone_shape(source_photo, slide)
    photo.left = x + Inches(0.12)
    photo.top = y + (row_height - photo_size) // 2
    photo.width = photo_size
    photo.height = photo_size

    text_left = x + Inches(0.65)
    text_width = column_width - Inches(0.75)
    box = slide.shapes.add_textbox(text_left, y, text_width, row_height)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    name_p = frame.paragraphs[0]
    name_p.space_after = Pt(0)
    name_run = name_p.add_run()
    name_run.text = name
    name_run.font.name = "CiscoSansTT"
    name_run.font.size = Pt(9)
    name_run.font.bold = True
    name_run.font.color.rgb = BLUE

    role_p = frame.add_paragraph()
    role_p.space_before = Pt(0)
    role_p.space_after = Pt(0)
    role_run = role_p.add_run()
    role_run.text = role
    role_run.font.name = "CiscoSansTT"
    role_run.font.size = Pt(7.5)
    role_run.font.color.rgb = GRAY


def add_team_column(slide, source_slide, x, leader_title):
    column_width = Inches(4.12)

    background = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x,
        Inches(0.98),
        column_width,
        Inches(6.15),
    )
    background.fill.solid()
    background.fill.fore_color.rgb = WHITE
    background.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEE)
    background.line.width = Pt(1)
    background.shadow.inherit = False

    heading = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x + Inches(0.08),
        Inches(1.06),
        column_width - Inches(0.16),
        Inches(0.34),
    )
    heading.fill.solid()
    heading.fill.fore_color.rgb = LIGHT_BLUE
    heading.line.fill.background()
    add_text(
        slide,
        leader_title,
        x + Inches(0.18),
        Inches(1.07),
        column_width - Inches(0.36),
        Inches(0.31),
        11,
        NAVY,
        bold=True,
    )

    leader_group = next(shape for shape in source_slide.shapes if shape.name == "Group 80")
    leader = clone_shape(leader_group, slide)
    leader.left = x + Inches(0.22)
    leader.top = Inches(1.48)
    leader.width = column_width - Inches(0.44)
    leader.height = Inches(0.90)

    reports = direct_report_data(source_slide)
    rows_top = Inches(2.48)
    rows_height = Inches(4.48)
    row_height = rows_height // len(reports)
    for index, report in enumerate(reports):
        add_report_row(
            slide,
            report,
            x,
            rows_top + index * row_height,
            column_width,
            row_height,
        )


def remove_original_slides(prs, slides_to_remove):
    for slide in slides_to_remove:
        # Find the slide id by its relationship target, then remove both references.
        for item in list(prs.slides._sldIdLst):
            if prs.part.related_part(item.rId) is slide.part:
                prs.part.drop_rel(item.rId)
                prs.slides._sldIdLst.remove(item)
                break


def main():
    prs = Presentation(SOURCE)
    source_slides = list(prs.slides)
    combined = prs.slides.add_slide(prs.slide_layouts[6])
    combined.background.fill.solid()
    combined.background.fill.fore_color.rgb = WHITE

    add_text(
        combined,
        "Americas Collaboration Partner Sales & Solution Engineering",
        Inches(0.38),
        Inches(0.20),
        Inches(12.6),
        Inches(0.56),
        25,
        NAVY,
        font_name="CiscoSansTT Medium",
    )
    add_text(
        combined,
        "Combined team overview",
        Inches(0.40),
        Inches(0.70),
        Inches(4.0),
        Inches(0.24),
        10,
        GRAY,
    )

    column_x = [Inches(0.30), Inches(4.61), Inches(8.92)]
    add_team_column(combined, source_slides[1], column_x[0], "Partner Sales · Chad Bullock")
    add_team_column(combined, source_slides[2], column_x[1], "Partner Sales · Kristeen Webb")
    add_team_column(
        combined,
        source_slides[3],
        column_x[2],
        "Solution Engineering · Stephanie Ross",
    )

    remove_original_slides(prs, source_slides)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
